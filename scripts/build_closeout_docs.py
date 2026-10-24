#!/usr/bin/env python3
"""Build the closeout package: one Markdown file and one accessible Word document.

    python scripts/build_closeout_docs.py

Sources (the only files anyone edits):
  Part 1  docs/closeout/CLIENTVERSE_CRM_CLOSEOUT_HANDOFF.md
  Part 2  docs/CLIENTVERSE_CRM_OWNER_ADMIN_MANUAL.md

Outputs (generated; never edit by hand, rebuild instead):
  docs/closeout/ClientVerse_CRM_Closeout_Package.md
  docs/closeout/ClientVerse_CRM_Closeout_Package.docx

The two parts are joined with every heading given a part-prefixed anchor (`p1-…`, `p2-…`),
so a link inside either part still lands on its own section, links between the parts
resolve inside the package, and links to other repository files point at GitHub.

The Word file is styled for reading: 13 pt body text, bold dark-navy headings, near-black
on white (contrast above 12:1), 1.3 line spacing with space between paragraphs, bordered
tables whose header row repeats on every page, a table of contents, and the document
language and title set for screen readers. Needs pandoc and python-docx.
"""

from __future__ import annotations

import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PART1 = ROOT / "docs/closeout/CLIENTVERSE_CRM_CLOSEOUT_HANDOFF.md"
PART2 = ROOT / "docs/CLIENTVERSE_CRM_OWNER_ADMIN_MANUAL.md"
OUT_MD = ROOT / "docs/closeout/ClientVerse_CRM_Closeout_Package.md"
OUT_DOCX = ROOT / "docs/closeout/ClientVerse_CRM_Closeout_Package.docx"
GITHUB = "https://github.com/ebyron357/Clientverse-crm/blob/main/"
TITLE = "ClientVerse CRM — Production Closeout Package"

NAVY = "1F3864"
TEXT = "1A1A1A"


def gfm_slug(text: str) -> str:
    """GitHub's heading anchor: lower case, punctuation dropped, spaces to hyphens."""
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)  # link text only
    text = text.replace("`", "").replace("*", "").strip().lower()
    text = re.sub(r"[^\w\- ]", "", text, flags=re.UNICODE)
    return text.replace(" ", "-")


def prepare(source: Path, prefix: str, for_docx: bool) -> str:
    """Give each heading a prefixed anchor and point every link at the package or GitHub."""
    lines = source.read_text(encoding="utf-8").splitlines()
    seen: dict[str, int] = {}
    mapping: dict[str, str] = {}
    out: list[str] = []
    fence = False
    for line in lines:
        if line.lstrip().startswith("```"):
            fence = not fence
        match = None if fence else re.match(r"^(#{1,6})\s+(.*?)\s*$", line)
        if match:
            slug = gfm_slug(match.group(2))
            count = seen.get(slug, 0)
            seen[slug] = count + 1
            original = slug if count == 0 else f"{slug}-{count}"
            anchor = f"{prefix}-{original}" if original else f"{prefix}-section-{len(mapping)}"
            mapping[original] = anchor
            if for_docx:
                line = f"{match.group(1)} {match.group(2)} {{#{anchor}}}"
            else:
                out.append(f'<a id="{anchor}"></a>')
                out.append("")
        out.append(line)
    text = "\n".join(out) + "\n"

    def internal(m: re.Match) -> str:
        target = m.group(1)
        return f"](#{mapping.get(target, prefix + '-' + target)})"

    text = re.sub(r"\]\(#([^)\s]+)\)", internal, text)

    base = source.parent

    def relative(m: re.Match) -> str:
        target, anchor = m.group(1), m.group(2) or ""
        resolved = (base / target).resolve()
        if resolved == PART2.resolve():
            return f"](#{('p2-' + anchor[1:]) if anchor else 'part-2'})"
        if resolved == PART1.resolve():
            return f"](#{('p1-' + anchor[1:]) if anchor else 'part-1'})"
        try:
            rel = resolved.relative_to(ROOT)
        except ValueError:
            return m.group(0)
        return f"]({GITHUB}{rel.as_posix()}{anchor})"

    return re.sub(r"\]\((?!https?:|mailto:|#)([^)#\s]+)(#[^)\s]*)?\)", relative, text)


def combined(for_docx: bool) -> str:
    part1 = prepare(PART1, "p1", for_docx)
    part2 = prepare(PART2, "p2", for_docx)
    if for_docx:
        head1, head2 = "# Part 1 — Production Closeout Handoff {#part-1}", "# Part 2 — Owner and Administrator Manual {#part-2}"
    else:
        head1 = '<a id="part-1"></a>\n\n# Part 1 — Production Closeout Handoff'
        head2 = '<a id="part-2"></a>\n\n# Part 2 — Owner and Administrator Manual'
    intro = (
        "This package is the complete owner and operator documentation for ClientVerse CRM. "
        "**Part 1** is the closeout handoff: what was delivered and verified, the production "
        "environment, access and recovery, configuration, operations, costs, test results, "
        "remaining owner actions and the closeout record. **Part 2** is the Owner and "
        "Administrator Manual: step-by-step instructions for every user and administrator "
        "workflow, with exact screen labels.\n\n"
        "It is generated from the two source documents in the repository "
        "(`docs/closeout/CLIENTVERSE_CRM_CLOSEOUT_HANDOFF.md` and "
        "`docs/CLIENTVERSE_CRM_OWNER_ADMIN_MANUAL.md`) by `scripts/build_closeout_docs.py`. "
        "To change it, edit a source and rebuild; this file is not edited by hand.\n"
    )
    # Each part's own H1 title becomes a subtitle under the part heading.
    demote = lambda text: re.sub(r"^# ", "## ", text, count=1, flags=re.M)  # noqa: E731
    return (f"{intro}\n{head1}\n\n{demote(part1)}\n\n{head2}\n\n{demote(part2)}")


def contents(text: str) -> str:
    """A linked list of the parts and their top sections; works in any viewer."""
    lines = ["## Contents {#contents}", ""]
    for match in re.finditer(r"^(#{1,3}) (.*?) \{#([^}]+)\}$", text, flags=re.M):
        level, title, anchor = len(match.group(1)), match.group(2), match.group(3)
        if level == 1:
            lines.append(f"- **[{title}](#{anchor})**")
        elif level == 2 and re.match(r"^(\d+\.|Appendix)", title):
            lines.append(f"    - [{re.sub(r'[*`]', '', title)}](#{anchor})")
    return "\n".join(lines) + "\n\n"


def build_reference(path: Path) -> None:
    """pandoc's default reference document, restyled for large, high-contrast reading."""
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt, RGBColor

    with open(path, "wb") as handle:
        handle.write(subprocess.run(["pandoc", "--print-default-data-file",
                                     "reference.docx"], check=True, capture_output=True).stdout)
    doc = Document(str(path))
    # Looked up by the name stored in the file: python-docx translates "Heading 1" to the
    # built-in "heading 1" on lookup, which pandoc's template does not use.
    by_name = {st.name: st for st in doc.styles}

    def style(name, size, bold=False, color=TEXT, before=0, after=8, font="Calibri"):
        st = by_name.get(name)
        if st is None:
            return
        st.font.name = font
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor.from_string(color)
        fmt = st.paragraph_format
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        fmt.line_spacing = 1.3

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        style(name, 13)
    style("Compact", 13, after=4)
    style("Title", 30, bold=True, color=NAVY, after=12)
    style("Subtitle", 18, color=NAVY)
    style("Heading 1", 24, bold=True, color=NAVY, before=24, after=12)
    style("Heading 2", 19, bold=True, color=NAVY, before=20, after=10)
    style("Heading 3", 16, bold=True, color=NAVY, before=16, after=8)
    style("Heading 4", 14, bold=True, color=NAVY, before=12, after=6)
    style("Heading 5", 13, bold=True, color=NAVY, before=10, after=6)
    style("Block Text", 13, color=TEXT)
    style("Source Code", 11, font="Consolas")
    if "Verbatim Char" in by_name:
        by_name["Verbatim Char"].font.name = "Consolas"
        by_name["Verbatim Char"].font.size = Pt(12)
    if "Hyperlink" in by_name:
        by_name["Hyperlink"].font.color.rgb = RGBColor.from_string("0B4F9C")
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Title"):
        if name in by_name:
            by_name[name].font.italic = False
            by_name[name].paragraph_format.keep_with_next = True
    doc.save(str(path))


def finish_docx(path: Path) -> None:
    """Bordered tables, a shaded header row that repeats, title and language set."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document(str(path))
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Owner and operator documentation"
    doc.core_properties.language = "en-US"
    for section in doc.sections:  # comfortable margins on Letter
        section.left_margin = section.right_margin = Pt(64)
        section.top_margin = section.bottom_margin = Pt(64)

    for table in list(doc.tables):
        # A key/value table written with an empty Markdown header row: drop that row.
        first = table.rows[0]
        if not any(cell.text.strip() for cell in first.cells):
            first._tr.getparent().remove(first._tr)
            header_row = False
        else:
            header_row = True
        tbl_pr = table._tbl.tblPr
        width = tbl_pr.find(qn("w:tblW"))
        if width is None:
            width = OxmlElement("w:tblW")
            tbl_pr.append(width)
        width.set(qn("w:w"), "5000")
        width.set(qn("w:type"), "pct")
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "8")
            element.set(qn("w:color"), "595959")
            borders.append(element)
        tbl_pr.append(borders)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "autofit")
        tbl_pr.append(layout)
        for row_index, row in enumerate(table.rows):
            is_header = header_row and row_index == 0
            if is_header:
                tr_pr = row._tr.get_or_add_trPr()
                header = OxmlElement("w:tblHeader")
                header.set(qn("w:val"), "true")
                tr_pr.append(header)
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "bottom", "left", "right"):
                    margin = OxmlElement(f"w:{edge}")
                    margin.set(qn("w:w"), "90")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                tc_pr.append(margins)
                if is_header:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:val"), "clear")
                    shade.set(qn("w:color"), "auto")
                    shade.set(qn("w:fill"), "DCE3EE")
                    tc_pr.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(12)
                        if is_header:
                            run.font.bold = True
    doc.save(str(path))


def main() -> int:
    OUT_MD.write_text(f"# {TITLE}\n\n" + combined(for_docx=False), encoding="utf-8")
    with tempfile.TemporaryDirectory() as tmp:
        reference = Path(tmp) / "reference.docx"
        build_reference(reference)
        source = Path(tmp) / "package.md"
        body = combined(for_docx=True)
        source.write_text(contents(body) + body, encoding="utf-8")
        subprocess.run([
            "pandoc", str(source), "-f", "gfm+attributes", "-t", "docx",
            "--reference-doc", str(reference),
            "--metadata", f"title={TITLE}",
            "--metadata", "subtitle=Closeout handoff and Owner/Administrator Manual",
            "--metadata", "lang=en-US",
            "-o", str(OUT_DOCX),
        ], check=True)
    finish_docx(OUT_DOCX)
    print(f"wrote {OUT_MD.relative_to(ROOT)} and {OUT_DOCX.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
